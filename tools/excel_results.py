"""Create and update public experiment result workbooks from eval.json files."""

from __future__ import annotations

import json
import os
import tempfile
from pathlib import Path
from typing import Iterable

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from tools.result_completion import is_complete_evaluation


METADATA_HEADERS = (
    "record_index",
    "repo_name",
    "repo_url",
    "test_framework",
    "production_file_path",
    "production_commit_SHA",
    "test_file_path",
    "test_commit_SHA",
    "commit_type",
    "changed_functions",
    "test_calls_changed_functions",
    "GroundTruth-Passing",
)

MODEL_METRICS = (
    "Status",
    "Error",
    "changed?",
    "CSR",
    "TPS",
    "UCR",
    "UCR-MatchDetail",
    "codeBLEU",
    "SPR",
    "Iterations",
    "WallTime",
    "Tokens",
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Arial", size=10)
BODY_FONT = Font(color="1F1F1F", name="Arial", size=10)
THIN_GRAY = Side(style="thin", color="D9E2F3")


def _display_value(value):
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def _record_row(record: dict, record_index: int) -> list:
    values = dict(record)
    values["record_index"] = record_index
    return [_display_value(values.get(header, "")) for header in METADATA_HEADERS]


def _header_map(worksheet) -> dict[str, int]:
    return {
        str(worksheet.cell(row=1, column=column).value): column
        for column in range(1, worksheet.max_column + 1)
        if worksheet.cell(row=1, column=column).value
    }


def _style_worksheet(worksheet) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.row_dimensions[1].height = 30

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=False)

    width_by_header = {
        "record_index": 13,
        "repo_name": 32,
        "repo_url": 42,
        "test_framework": 16,
        "production_file_path": 48,
        "production_commit_SHA": 42,
        "test_file_path": 48,
        "test_commit_SHA": 42,
        "commit_type": 16,
        "changed_functions": 36,
        "test_calls_changed_functions": 26,
        "GroundTruth-Passing": 22,
    }
    for header, column in _header_map(worksheet).items():
        width = width_by_header.get(header, 16)
        if header.endswith("-Error"):
            width = 42
        elif header.endswith("-UCR-MatchDetail"):
            width = 38
        worksheet.column_dimensions[get_column_letter(column)].width = width


def _ensure_sheet(workbook, method: str, records: list[dict]):
    if method in workbook.sheetnames:
        worksheet = workbook[method]
    else:
        worksheet = workbook.create_sheet(method)
        worksheet.append(list(METADATA_HEADERS))

    headers = _header_map(worksheet)
    for header in METADATA_HEADERS:
        if header not in headers:
            worksheet.cell(row=1, column=worksheet.max_column + 1, value=header)
            headers = _header_map(worksheet)

    existing_rows: dict[int, int] = {}
    index_column = headers["record_index"]
    for row_number in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row=row_number, column=index_column).value
        if isinstance(value, int):
            existing_rows[value] = row_number

    for record_index, record in enumerate(records):
        row_number = existing_rows.get(record_index)
        row_values = _record_row(record, record_index)
        if row_number is None:
            worksheet.append(row_values)
            continue
        for column, value in enumerate(row_values, start=1):
            worksheet.cell(row=row_number, column=column, value=value)

    _style_worksheet(worksheet)
    return worksheet


def ensure_result_workbook(
    excel_path: Path,
    records: list[dict],
    methods: Iterable[str],
) -> None:
    """Create a workbook or synchronize its sample rows and method sheets."""
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    if excel_path.exists():
        workbook = openpyxl.load_workbook(excel_path)
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)

    for method in methods:
        _ensure_sheet(workbook, method, records)
    _save_atomic(workbook, excel_path)


def _ensure_model_columns(worksheet, model: str) -> dict[str, int]:
    headers = _header_map(worksheet)
    columns: dict[str, int] = {}
    for metric in MODEL_METRICS:
        header = f"{model}-{metric}"
        column = headers.get(header)
        if column is None:
            column = worksheet.max_column + 1
            worksheet.cell(row=1, column=column, value=header)
            headers[header] = column
        columns[metric] = column
    _style_worksheet(worksheet)
    return columns


def _find_sample_row(worksheet, record: dict, record_index: int) -> int:
    headers = _header_map(worksheet)
    expected = {
        "record_index": record_index,
        "repo_name": record.get("repo_name", ""),
        "test_commit_SHA": record.get("test_commit_SHA", ""),
        "test_file_path": record.get("test_file_path", ""),
    }
    for row_number in range(2, worksheet.max_row + 1):
        if all(
            worksheet.cell(row=row_number, column=headers[field]).value == value
            for field, value in expected.items()
        ):
            return row_number
    raise ValueError(
        f"Workbook sheet {worksheet.title!r} does not contain record {record_index}"
    )


def _ucr_detail(case_comparison: dict) -> str:
    summary = case_comparison.get("match_summary") or {}
    return (
        f"exact={int(summary.get('exact', 0) or 0)}; "
        f"normalized={int(summary.get('normalized', 0) or 0)}; "
        f"semantic={int(summary.get('semantic', 0) or 0)}; "
        f"llm={int(case_comparison.get('llm_passing_count', 0) or 0)}; "
        "ground_truth="
        f"{int(case_comparison.get('ground_truth_passing_count', 0) or 0)}"
    )


def apply_evaluation(
    excel_path: Path,
    records: list[dict],
    evaluation: dict,
) -> None:
    """Write one completed evaluation to its method sheet and sample row."""
    method = str(evaluation.get("mode", ""))
    model = str(evaluation.get("llm", ""))
    record_index = int(evaluation["record_index"])
    record = records[record_index]

    workbook = openpyxl.load_workbook(excel_path)
    if method not in workbook.sheetnames:
        _ensure_sheet(workbook, method, records)
    worksheet = workbook[method]
    row_number = _find_sample_row(worksheet, record, record_index)
    columns = _ensure_model_columns(worksheet, model)

    ground_truth = evaluation.get("ground_truth_test_result") or {}
    ground_truth_passing = (
        int(ground_truth.get("passing_count", 0) or 0)
        if ground_truth.get("success") is True
        else None
    )
    for target in workbook.worksheets:
        target_headers = _header_map(target)
        if "GroundTruth-Passing" not in target_headers:
            continue
        target_row = _find_sample_row(target, record, record_index)
        target.cell(
            row=target_row,
            column=target_headers["GroundTruth-Passing"],
            value=ground_truth_passing,
        )

    comparison = evaluation.get("comparison") or {}
    run_result = evaluation.get("llm_run_result") or {}
    case_comparison = evaluation.get("case_comparison") or {}
    test_result = run_result.get("test_result")
    if test_result is None:
        test_result = "pass" if evaluation.get("test_passed") else "fail"

    values = {
        "Status": "completed",
        "Error": "",
        "changed?": evaluation.get("llm_test_changed", ""),
        "CSR": 1 if run_result.get("compilation") == "pass" else 0,
        "TPS": 1 if test_result == "pass" else 0,
        "UCR": case_comparison.get("ucr"),
        "UCR-MatchDetail": _ucr_detail(case_comparison),
        "codeBLEU": comparison.get("similarity"),
        "SPR": 1 if evaluation.get("syntax_pass") else 0,
        "Iterations": evaluation.get("iterations"),
        "WallTime": evaluation.get("wall_time"),
        "Tokens": evaluation.get("tokens"),
    }
    for metric, value in values.items():
        worksheet.cell(row=row_number, column=columns[metric], value=value)

    worksheet.cell(row=row_number, column=columns["UCR"]).number_format = "0.0"
    worksheet.cell(row=row_number, column=columns["codeBLEU"]).number_format = "0.0000"
    worksheet.cell(row=row_number, column=columns["WallTime"]).number_format = "0.00"
    _save_atomic(workbook, excel_path)


def apply_job_failure(
    excel_path: Path,
    records: list[dict],
    *,
    record_index: int,
    method: str,
    model: str,
    message: str,
) -> None:
    """Record a failed Docker job without inventing metric values."""
    workbook = openpyxl.load_workbook(excel_path)
    worksheet = workbook[method]
    row_number = _find_sample_row(worksheet, records[record_index], record_index)
    columns = _ensure_model_columns(worksheet, model)
    for metric in MODEL_METRICS:
        worksheet.cell(row=row_number, column=columns[metric]).value = None
    worksheet.cell(row=row_number, column=columns["Status"], value="failed")
    worksheet.cell(row=row_number, column=columns["Error"], value=message[:2000])
    _save_atomic(workbook, excel_path)


def find_evaluation(
    output_dir: Path,
    record: dict,
    *,
    record_index: int,
    method: str,
    model: str,
    modified_after_ns: int | None = None,
) -> dict | None:
    """Find the completed evaluation matching one exact requested job."""
    method_dir = output_dir / f"output_{method}"
    if not method_dir.is_dir():
        return None
    for eval_path in sorted(method_dir.glob("*/*/*/eval.json")):
        if modified_after_ns is not None:
            try:
                if eval_path.stat().st_mtime_ns < modified_after_ns:
                    continue
            except OSError:
                continue
        try:
            evaluation = json.loads(eval_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        metadata = evaluation.get("metadata") or {}
        if (
            is_complete_evaluation(evaluation)
            and evaluation.get("record_index") == record_index
            and evaluation.get("mode") == method
            and evaluation.get("llm") == model
            and metadata.get("repo_name", "") == record.get("repo_name", "")
            and metadata.get("test_file_path", "") == record.get("test_file_path", "")
            and metadata.get("test_commit_SHA", "") == record.get("test_commit_SHA", "")
        ):
            return evaluation
    return None


def aggregate_evaluations(
    excel_path: Path,
    output_dir: Path,
    records: list[dict],
    *,
    record_index: int,
    methods: Iterable[str],
    models: Iterable[str],
) -> int:
    """Write all matching existing evaluations for selected jobs to Excel."""
    updated = 0
    record = records[record_index]
    for method in methods:
        for model in models:
            evaluation = find_evaluation(
                output_dir,
                record,
                record_index=record_index,
                method=method,
                model=model,
            )
            if evaluation is None:
                continue
            apply_evaluation(excel_path, records, evaluation)
            updated += 1
    return updated


def _save_atomic(workbook, excel_path: Path) -> None:
    """Save through a sibling temporary file so interruption preserves the prior copy."""
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{excel_path.stem}.", suffix=".xlsx", dir=excel_path.parent
    )
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, excel_path)
    finally:
        temporary_path.unlink(missing_ok=True)
